#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
"""Generate docs/SANKHYA-deck.pptx, the jury deck, from the report's slot values.

Called by tools/make_report.py with the `F` dictionary it computed from the CSVs in
bench/results/, so every number on a slide is the one docs/REPORT.md prints, and each slide
names the CSV it reads. The prose lives here; a number can only be a `{slot}`. The file is
byte-for-byte reproducible (fixed core properties from python-pptx's template, and every zip
entry re-stamped to a fixed date), which is what lets CI compare the committed deck with a
fresh one.

    python tools/make_deck.py      # docs/SANKHYA-deck.pptx from the CSVs on main
"""
from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "tools"))

NAVY = RGBColor(0x14, 0x2B, 0x4F)
ORANGE = RGBColor(0xD9, 0x6C, 0x1F)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
W, H = Inches(13.333), Inches(7.5)
SLIDES = 14

TEAM = [("Bhumika Gurav", "team lead"), ("Deekshith K", ""), ("Anjankumar L N", ""),
        ("Abhijeet Sharma", ""), ("Ayush V Upadhye", ""), ("Chirag Honnyal", "")]


def plain(text: str) -> str:
    return re.sub(r"[`*]", "", str(text))


def cells(markdown_table: str) -> list[list[str]]:
    """The first Markdown pipe table in the text as rows of plain cells, the separator row
    left out (the head-to-head slots carry a second, per-instance table after a blank)."""
    rows = []
    for line in str(markdown_table).strip().split("\n\n")[0].splitlines():
        if line.startswith("|") and not re.match(r"^\|\s*:?-", line):
            rows.append([plain(c.strip()) for c in line.strip().strip("|").split("|")])
    return rows


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.blank = self.prs.slide_layouts[6]

    def text(self, slide, x, y, w, h, lines, size=16, color=GREY, bold=False, font=None):
        box = slide.shapes.add_textbox(x, y, w, h)
        frame = box.text_frame
        frame.word_wrap = True
        for index, line in enumerate(lines if isinstance(lines, list) else [lines]):
            para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = para.add_run()
            run.text = plain(line)
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
            if font:
                run.font.name = font
            para.space_after = Pt(size * 0.45)
        return box

    def slide(self, title: str, source: str) -> object:
        slide = self.prs.slides.add_slide(self.blank)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(1.2))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        self.text(slide, Inches(0.5), Inches(0.15), W - Inches(1), Inches(1), title, 24,
                  WHITE, True)
        foot = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, H - Inches(0.45), W, Inches(0.45))
        foot.fill.solid()
        foot.fill.fore_color.rgb = LIGHT
        foot.line.fill.background()
        self.text(slide, Inches(0.5), H - Inches(0.42), W - Inches(1), Inches(0.4),
                  f"Source: {source}", 10, GREY)
        return slide

    def table(self, slide, rows, x, y, w, size=12, first_bold=True):
        shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w,
                                       Inches(0.35) * len(rows))
        table = shape.table
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.text = plain(value)
                for para in cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.size = Pt(size)
                        run.font.bold = first_bold and r == 0
                        run.font.color.rgb = WHITE if r == 0 else GREY
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY if r == 0 else (WHITE if r % 2 else LIGHT)
        return shape

    def bars(self, slide, items, x, y, w, h, unit=" s"):
        """Horizontal bars, one per (label, value), the longest filling the width."""
        top = max(v for _, v in items) or 1.0
        row = h / len(items)
        for index, (label, value) in enumerate(items):
            yy = y + int(row * index)
            self.text(slide, x, yy, Inches(1.6), row, label, 12, GREY)
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x + Inches(1.7), yy + Emu(row // 5),
                                         int((w - Inches(3.2)) * value / top) or Emu(1),
                                         int(row * 0.6))
            bar.fill.solid()
            bar.fill.fore_color.rgb = ORANGE if label == "SANKHYA" else NAVY
            bar.line.fill.background()
            self.text(slide, bar.left + bar.width + Inches(0.1), yy, Inches(1.5), row,
                      f"{value:g}{unit}", 12, GREY)

    def big(self, slide, x, y, number, caption, w=Inches(3)):
        self.text(slide, x, y, w, Inches(0.9), number, 36, ORANGE, True)
        self.text(slide, x, y + Inches(0.85), w, Inches(1.2), caption, 12, GREY)

    def save(self, target: Path) -> None:
        raw = io.BytesIO()
        self.prs.save(raw)
        # Stored, not deflated: the bytes then depend only on the XML, not on the zlib the
        # interpreter links, so Windows and the CI runner write the same file.
        with zipfile.ZipFile(raw) as src, zipfile.ZipFile(target, "w", zipfile.ZIP_STORED) as dst:
            for info in src.infolist():
                stamped = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
                dst.writestr(stamped, src.read(info.filename))


def sgm_items(table_md: str) -> list[tuple[str, float]]:
    rows = cells(table_md)
    out = []
    for row in rows[1:]:
        if row[-1].endswith(" s"):
            out.append((row[0], float(row[-1][:-2])))
    return out


def write(F: dict, target: Path) -> None:  # noqa: C901 - one slide after another
    f = lambda text: text.format_map(F)  # noqa: E731
    d = Deck()

    # 1 title
    s = d.prs.slides.add_slide(d.blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    d.text(s, Inches(0.8), Inches(1.6), W - Inches(1.6), Inches(1.2), "SANKHYA", 60, WHITE, True)
    d.text(s, Inches(0.8), Inches(2.8), W - Inches(1.6), Inches(1.6),
           ["An indigenous optimization solver: LP, MILP, convex QP and NLP, written from the "
            "textbooks in C++20, with a GPU first-order engine and an independent checker on "
            "every answer.",
            "Smart India Hackathon 2026 - problem statement SIH26119 - Mangalore Refinery and "
            "Petrochemicals Limited"], 20, WHITE)
    d.text(s, Inches(0.8), H - Inches(1.4), W - Inches(1.6), Inches(1),
           ["Every number in this deck is computed from a CSV in bench/results/ on main, named "
            "on the slide, by tools/make_deck.py; the same CSVs produce docs/REPORT.pdf and "
            "docs/BENCHMARKS.md. github.com/thegoodengineers/SANKHYA - sankhya-solver.vercel.app"],
           12, LIGHT)

    # 2 problem
    s = d.slide("The problem: a refinery plans in LPs, MILPs and QPs, on someone else's solver",
                "the problem statement SIH26119 as issued; docs/PS26119_COVERAGE.md")
    d.text(s, Inches(0.5), Inches(1.4), Inches(6.2), Inches(5), [
        "Which crudes to buy, how hard to run each unit, how to blend to specification: the "
        "weekly plan is a linear program, the unit on/off and minimum-run decisions make it "
        "mixed-integer, and price-dependent blending makes it quadratic.",
        "The problem statement asks for an indigenous, GPU-accelerated solver as a sovereign "
        "alternative to CPLEX and Xpress: LP, MILP, QP, presolve, cuts, sensitivity, multi-core "
        "and GPU, benchmarked on Netlib, MIPLIB and Mittelmann, and not built on any existing "
        "solver.",
        "SANKHYA is that solver, from a blank repository on 25 August 2026. What follows is "
        "what is on main today, with the file behind each number."], 15)
    d.text(s, Inches(7.2), Inches(1.4), Inches(5.6), Inches(5), [
        "What a judge can do in five minutes",
        "1. pip install the wheel from the release, or unpack an archive: no compiler.",
        "2. sankhya solve demo/crude_blend.mps --write-sol blend.sol",
        "3. python tools/verify_solution.py demo/crude_blend.mps blend.sol: a checker that "
        "shares no code with the solver re-derives the answer.",
        "4. demo/finale.sh: one refinery MILP solved, proved in exact arithmetic, priced, "
        "broken and repaired, re-planned from yesterday's basis.",
        "5. Open docs/REPORT.pdf: every claim with its CSV and commit."], 14, NAVY)

    # 3 architecture
    s = d.slide("Architecture: one Model, one solve(), several engines, one checker outside",
                "docs/ARCHITECTURE.md; docs/PROVENANCE.md section 3 for the paper above each file")
    columns = [
        ("Read", ["MPS and LP readers (src/io)", "gz inputs", "Model: the frozen interface "
                  "every reader produces and every engine consumes"]),
        ("Presolve", ["reductions with a postsolve stack (src/presolve)", "scaling by powers of "
                      "two, geometric prescaling on retry", "convexity test for a QP: LDL^T, a "
                      "negative pivot is the certificate"]),
        ("Solve", ["revised primal and dual simplex: Markowitz LU, Forrest-Tomlin, dual "
                   "steepest edge (src/simplex, src/la)", "restarted PDHG on the CPU and, "
                   "through CUDA, on the card (src/pdhg, src/gpu)", "Mehrotra interior point, "
                   "sparse LDL^T, crossover (src/ipm)", "proximal interior point for convex QP "
                   "(src/qp)", "branch and cut: cuts, heuristics, SOS, certificates (src/mip)",
                   "local NLP and convex MINLP (src/nlp)"]),
        ("Report", ["Solution: objective, point, duals, ranging, basis, certificate",
                    "C API, CLI, ctypes Python bindings, PuLP / Pyomo / CVXPY adapters",
                    "tools/verify_solution.py and tools/verify_certificate.py: read only the "
                    "files, never link the solver"]),
    ]
    width = (W - Inches(1)) / len(columns)
    for index, (head, items) in enumerate(columns):
        x = Inches(0.5) + int(width * index)
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(1.4), int(width) - Inches(0.15),
                                 Inches(0.5))
        box.fill.solid()
        box.fill.fore_color.rgb = ORANGE
        box.line.fill.background()
        d.text(s, x, Inches(1.45), int(width) - Inches(0.15), Inches(0.5), head, 16, WHITE,
               True)
        d.text(s, x, Inches(2.0), int(width) - Inches(0.2), Inches(4.8),
               ["- " + item for item in items], 12)

    # 4 correctness
    s = d.slide("Correctness evidence: an answer is reported only after a separate program agrees",
                f("{netlib_csv}, {infeasible_csv}, {certified_csv}, {exact_csv}, {stress_csv}"))
    d.table(s, [
        ["what is checked", "result", "CSV", "commit"],
        ["Netlib LP, full set: accepted by the independent verifier",
         f("{netlib_verified} of {netlib_n}"), f("{netlib_csv}"), f("{netlib_commit}")],
        ["Netlib LP: within 1e-6 of Koch's exact rational optimum and verified",
         f("{netlib_exact} of {netlib_n}"), f("{netlib_csv}"), f("{netlib_commit}")],
        ["Netlib infeasible LPs: verdict with a Farkas certificate the verifier accepts",
         f("{infeasible_passed} of {infeasible_n}"), f("{infeasible_csv}"),
         f("{infeasible_commit}")],
        ["Optimal Netlib answers with a rigorous (Neumaier-Shcherbina) bound within 1e-6",
         f("{certified_tight} of {certified_optimal}"), f("{certified_csv}"),
         f("{certified_commit}")],
        ["Duals, reduced costs and ranges re-derived in rational arithmetic (LPs done)",
         f("{exact_done} of {exact_n}; {exact_corrected} of {exact_values} values corrected"),
         f("{exact_csv}"), f("{exact_commit}")],
        ["Stress set of badly scaled and adversarial LPs: correct / wrong / declined",
         f("{stress_sankhya_correct} / {stress_sankhya_wrong} / {stress_sankhya_failed} of "
           "{stress_sankhya_n}"), f("{stress_csv}"), f("{stress_commit}")],
        ["HiGHS {stress_highs_version} on its defaults, same set: correct / wrong / failed".format_map(F),
         f("{stress_highs_correct} / {stress_highs_wrong} / {stress_highs_failed} of "
           "{stress_highs_n}"), f("{stress_csv}"), f("{stress_commit}")],
    ], Inches(0.5), Inches(1.4), W - Inches(1), 11)
    d.text(s, Inches(0.5), Inches(5.6), W - Inches(1), Inches(1.2), [
        "tools/verify_solution.py reads the model file and the written .sol file with its own "
        "parsers and never links the C++. A verdict of infeasible or unbounded arrives with a "
        "certificate it checks the same way; a MILP's proof is a VIPR certificate checked in "
        "exact arithmetic (tools/verify_certificate.py)."], 13)

    # 5 head-to-head LP
    s = d.slide("Head-to-head on LP: HiGHS, SCIP, CBC/Clp and GLPK, same machine, one thread",
                f("{h2h_netlib_csv} ({h2h_netlib_commit}); {h2h_kennington_csv} "
                  "({h2h_kennington_commit}); method in docs/BENCHMARKS.md section 4a"))
    d.text(s, Inches(0.5), Inches(1.3), Inches(6), Inches(0.5),
           f("Netlib LP, {h2h_netlib_n} instances, {h2h_netlib_limit} s limit: shifted "
             "geometric mean time, shift 10 s"), 14, NAVY, True)
    d.bars(s, sgm_items(F["h2h_netlib_table"]), Inches(0.5), Inches(1.9), Inches(6), Inches(2.4))
    d.text(s, Inches(6.9), Inches(1.3), Inches(6), Inches(0.5),
           f("Kennington LP, {h2h_kennington_n} instances, {h2h_kennington_limit} s limit"),
           14, NAVY, True)
    d.bars(s, sgm_items(F["h2h_kennington_table"]), Inches(6.9), Inches(1.9), Inches(6),
           Inches(2.4))
    d.text(s, Inches(0.5), Inches(4.5), W - Inches(1), Inches(2.4), [
        f("Against HiGHS instance by instance: Netlib faster on {h2h_netlib_wins}, tie on "
          "{h2h_netlib_ties}, slower on {h2h_netlib_losses} of {h2h_netlib_compared}; "
          "Kennington faster on {h2h_kennington_wins}, tie on {h2h_kennington_ties}, slower on "
          "{h2h_kennington_losses} of {h2h_kennington_compared} (a tie is within 10%)."),
        f("The same verifier checks every rival's answer; on Netlib it rejected "
          "{h2h_netlib_rejected} rival answers. A run that is not optimal, at the published "
          "optimum and verified is charged the full time limit (Mittelmann's convention)."),
        "SANKHYA is not the fastest LP solver here and this deck does not say it is: it is the "
        "one whose every answer, and every rival's, went through the same independent check."],
        13)

    # 6 QP and MILP
    s = d.slide("Convex QP and MILP: where the checker changes the ranking",
                f("{h2h_maros_csv} ({h2h_maros_commit}); {maros_csv} ({maros_commit}); "
                  "{miplib_csv} ({miplib_commit})"))
    rows = [r[:1] + r[2:6] + r[-1:] for r in cells(F["h2h_maros_table"])]
    rows[0] = ["solver", "runs", "solved", "matched", "verified", "SGM time, shift 10 s"]
    d.text(s, Inches(0.5), Inches(1.3), W - Inches(1), Inches(0.5),
           f("Maros-Meszaros convex QP, {h2h_maros_n} instances, {h2h_maros_limit} s limit, "
             "head-to-head"), 14, NAVY, True)
    d.table(s, rows, Inches(0.5), Inches(1.8), Inches(7.5), 12)
    d.text(s, Inches(8.3), Inches(1.8), Inches(4.6), Inches(3), [
        f("A rival's optimal status is not taken on trust either: the verifier rejected "
          "{h2h_maros_rejected} rival QP answers in this run, and only verified runs count "
          "for time. Against HiGHS instance by instance: faster on {h2h_maros_wins}, tie on "
          "{h2h_maros_ties}, slower on {h2h_maros_losses} of {h2h_maros_compared}."),
        f("Alone, on the laptop: {maros_matched} of {maros_n} at the published objective, "
          "{maros_rejected} optimal answers rejected ({maros_csv}).")], 12)
    d.big(s, Inches(0.5), Inches(4.9), f("{miplib_matched} of {miplib_n}"),
          f("MIPLIB 2017 easy set at the published optimum, {miplib_proved} of {miplib_n} "
            "proved optimal inside the limit ({miplib_csv})"), Inches(4))
    d.big(s, Inches(4.8), Inches(4.9), f("{hs_matched} of {hs_n}"),
          f("Hock-Schittkowski NLP at the published objective, {hs_verified} accepted by the "
            "checker ({hs_csv})"), Inches(4))
    d.big(s, Inches(9.1), Inches(4.9), f("{minlp_matched} of {minlp_n}"),
          f("convex MINLPLib at the published objective, {minlp_verified} accepted "
            "({minlp_csv})"), Inches(4))

    # 7 GPU
    s = d.slide(f("GPU: the first-order engine on an {gpu_gpu}, against the faster CPU arm"),
                f("{gpu_csv} ({gpu_commit}); docs/BENCHMARKS.md section 1g.3; the L4 run beside "
                  "it in gpu-real-l4-58a8374.csv"))
    rows = cells(F["gpu_table"])
    d.table(s, rows, Inches(0.5), Inches(1.4), W - Inches(1), 12)
    d.text(s, Inches(0.5), Inches(1.6) + Inches(0.35) * len(rows) + Inches(0.3), W - Inches(1),
           Inches(2.6), [
        f("{gpu_worst} to {gpu_best} over {gpu_cells} model-tolerance cells, {gpu_limits} of "
          "them ending at the time limit on every arm. The card wins where the model is large "
          "and loses below about 10,000 rows, where launching kernels costs more than the "
          "arithmetic; the readme and section 1g name the losses."),
        "The CUDA port is the same restarted PDHG: cuSPARSE products, CUB reductions in a "
        "fixed order (deterministic), the iteration loop captured as a CUDA graph, convergence "
        "evaluated on the device. A CPU build with no card warns and solves on the CPU; the "
        "CUDA archive is compiled in CI without a card.",
        "HiGHS on the same host is still faster on brazil3 than the card at 1e-8 (readme, "
        "section 1g.3): the GPU story is a real speed-up of our own engine, not a claim over "
        "the field."], 13)

    # 8 scale
    s = d.slide("Scale: a million rows on the CPU, and where large LPs stop today",
                f("{million_csv} ({million_commit}); {mittelmann_csv} ({mittelmann_commit}); "
                  "{kennington_csv} ({kennington_commit})"))
    d.big(s, Inches(0.5), Inches(1.5), f("{million_done} of {million_n}"),
          f("engine runs at a million rows optimal and verified: {million_list}"), Inches(6))
    d.big(s, Inches(7), Inches(1.5), f("{kennington_passed} of {kennington_n}"),
          f("Kennington LPs matched to the published optimum and verified ({kennington_csv})"),
          Inches(5.5))
    d.big(s, Inches(0.5), Inches(4.2), f("{mittelmann_passed} of {mittelmann_n}"),
          f("Mittelmann LPs solved and verified inside the limit ({mittelmann_csv}). This is "
            "where the solver stops: the rest end on the interior point's ordering budget or "
            "the time limit, named row by row in docs/BENCHMARKS.md section 1d."), Inches(6))
    d.text(s, Inches(7), Inches(4.2), Inches(5.5), Inches(2.5), [
        "Large LPs are routed by the normal-equations bound and the dual's work: the simplex "
        "below, the interior point above, PDHG for what neither factors. Time limits reach "
        "inside the factorization, so a run stops on time, not when the factor happens to "
        "finish."], 13)

    # 9 refinery demo
    s = d.slide("The refinery demo: one MILP solved, proved, priced, broken, repaired, re-planned",
                f("docs/FINALE.md; {scenarios_csv} ({scenarios_commit}); {million_csv} "
                  "({million_commit}); model from bench/case_studies/refinery/generator.py"))
    d.text(s, Inches(0.5), Inches(1.4), Inches(7.4), Inches(5.2), [
        "demo/finale.sh, eight steps, each one printed line and its time, the script stopping "
        "at the first check that fails:",
        "1. solve the small refinery MILP with a VIPR certificate (on the card when there is one)",
        "2. prove it: the certificate in exact rational arithmetic, the plan by the verifier",
        "3. plan: the binding limits and their shadow prices, each re-solved one unit looser, "
        "so a degenerate price is not mistaken for the value of the next unit",
        "4. repair: a delivery commitment raised until the plant cannot meet it; the Farkas "
        "proof checked, the smallest repair found",
        "5. bundle: an evidence bundle made and replayed, the proof checked again",
        "6. prices: twenty crude price sets through sankhya scenarios",
        "7. replan: today's prices and demands, cold and from yesterday's basis (--warm-start)",
        "8. units: a minimum run rate as semi-continuous columns, branched natively and through "
        "the binary reformulation, the two answers verified and equal",
        "Synthetic data from a fixed seed: no real industrial model has been run (a stated "
        "limit)."], 12)
    d.big(s, Inches(8.3), Inches(1.5), f("{scenarios_verified} of {scenarios_n}"),
          f("scenarios verified, each re-solved from the base basis by the dual simplex and "
            "checked against its own model, agreeing with its single solve ({scenarios_csv})"),
          Inches(4.6))
    d.big(s, Inches(8.3), Inches(4.3), f("{million_largest}"),
          f("the refinery family at a million rows, optimal and verified on the CPU "
            "({million_list})"), Inches(4.6))

    # 10 packaging
    s = d.slide("Packaging: no compiler needed, checked on a fresh machine by CI",
                ".github/workflows/release.yml; docs/PROVENANCE.md section 2b; the v1.0.0-rc2 "
                "release")
    d.text(s, Inches(0.5), Inches(1.4), Inches(6.2), Inches(5.2), [
        "Release archives: Linux x86-64 CPU and CUDA (runtime linked in), Windows x64 CPU, "
        "each unpacked and run on a runner with no compiler before the draft release exists.",
        "Python wheels: ctypes over the C API, one wheel per platform for CPython 3.9 to 3.13, "
        "each installed into a clean venv of every version, the demo solved and verified there.",
        "Container image on ghcr.io: the CUDA build, the verifier and the demo models.",
        "C API (include/sankhya/sankhya.h), CLI (solve, diagnose, scenarios, engines, info), "
        "Python bindings, and adapters that make SANKHYA the solver under PuLP, Pyomo and "
        "CVXPY models.",
        "Every archive's MANIFEST.txt records the commit, compiler and libraries; CI fails the "
        "packaging if any of them is an optimization solver."], 13)
    d.text(s, Inches(7), Inches(1.5), Inches(5.8), Inches(3), [
        "pip install https://github.com/thegoodengineers/SANKHYA/releases/download/v1.0.0-rc2/"
        "sankhya-1.0.0-py3-none-manylinux_2_35_x86_64.whl",
        "python -c \"import sankhya; print(sankhya.Model.read('demo/crude_blend.mps')"
        ".solve().objective)\"",
        "sankhya solve demo/crude_blend.mps --write-sol blend.sol",
        "python tools/verify_solution.py demo/crude_blend.mps blend.sol"], 11, NAVY,
        font="Courier New")

    # 11 limits
    s = d.slide("Honest limits: what the numbers do not say",
                f("{h2h_netlib_csv}; {h2h_kennington_csv}; {mittelmann_csv}; {miplib_csv}; "
                  "{stress_csv}; {certified_csv}; docs/NEGATIVE-RESULTS.md"))
    d.text(s, Inches(0.5), Inches(1.4), W - Inches(1), Inches(5.4), [
        f("Speed against HiGHS: Netlib SGM {h2h_netlib_sankhya_sgm} s against "
          "{h2h_netlib_highs_sgm} s; Kennington {h2h_kennington_sankhya_sgm} s against "
          "{h2h_kennington_highs_sgm} s, slower on {h2h_kennington_losses} of "
          "{h2h_kennington_compared} instances."),
        f("Large LPs: Mittelmann {mittelmann_passed} of {mittelmann_n} inside the limit."),
        f("MILP proofs: {miplib_proved} of {miplib_n} MIPLIB 2017 instances proved, "
          "{miplib_matched} at the published optimum."),
        f("Declined instances: stress set {stress_sankhya_failed} of {stress_sankhya_n} "
          "declined ({stress_sankhya_wrong} wrong); Maros-Meszaros {maros_optimal} of "
          "{maros_n} optimal; safe bounds, {certified_none} of {certified_optimal} optimal "
          "answers with no finite bound, and they say so."),
        f("The GPU: {gpu_worst} at the low end of the {gpu_cells} cells, a loss below about "
          "10,000 rows, {gpu_limits} cells at the time limit; HiGHS on the same host beats the "
          "card on brazil3 at 1e-8."),
        "No real industrial model: every refinery number is from a generated model with a fixed "
        "seed.",
        "What was built and then demoted, left opt-in or withdrawn, with the number that did "
        "it, is kept by name in docs/NEGATIVE-RESULTS.md."], 14)

    # 12 provenance
    s = d.slide("Provenance: not built on any solver, and CI proves it on every build",
                "docs/PROVENANCE.md; docs/sbom.spdx.json; .github/workflows/ci.yml, job provenance")
    d.text(s, Inches(0.5), Inches(1.4), W - Inches(1), Inches(5.4), [
        "The red line: no source from CBC, Clp, HiGHS, SCIP, SoPlex, GLPK, lp_solve, OSQP, PDLP, "
        "cuPDLP, OR-Tools, Gurobi, CPLEX or Xpress was copied, vendored, linked or read: not a "
        "simplex, not a cut, not an MPS reader.",
        "Every algorithm names the paper or textbook it was written from, above the code and in "
        "the citation table (Dantzig, Chvatal, Maros, Koberstein, Forrest-Tomlin, Harris, "
        "Markowitz, Chambolle-Pock, Applegate et al., Mehrotra, Neumaier-Shcherbina, ...).",
        "Dependencies are the non-solver kind, each with its licence and one line on why it is "
        "not a solver: fmt, CLI11, nlohmann/json, zlib, GoogleTest, cuSPARSE, CUB, OpenMP; "
        "HiGHS, SCIP, CBC and GLPK exist only as separate processes in the benchmark harness.",
        "CI prints the full link line and ldd of the binary and fails if a solver library ever "
        "appears; the SPDX SBOM is committed, not only produced; every path the citation table "
        "names must exist.",
        "Judgement calls, where a line could be argued, are logged with their reasoning in the "
        "same document rather than decided silently."], 14)

    # 13 team
    s = d.slide("Team: the good engineers", "sankhya-solver.vercel.app/team; git log")
    for index, (name, role) in enumerate(TEAM):
        x = Inches(0.5) + int((W - Inches(1)) / 3 * (index % 3))
        y = Inches(1.6) + Inches(2.2) * (index // 3)
        d.text(s, x, y, Inches(4), Inches(0.6), name, 22, NAVY, True)
        d.text(s, x, y + Inches(0.6), Inches(4), Inches(0.5), role or "engineer", 14)
    d.text(s, Inches(0.5), Inches(6.1), W - Inches(1), Inches(0.8), [
        "Rules the repository enforces on us: no number without a CSV, correct over fast, the "
        "paper cited above the code, and what is missing said with the number."], 13)

    # 14 close
    s = d.slide("Where to look", "README.md; docs/REPORT.pdf; docs/BENCHMARKS.md; "
                "docs/VIDEO.md; sankhya-solver.vercel.app")
    d.text(s, Inches(0.5), Inches(1.5), W - Inches(1), Inches(5), [
        "github.com/thegoodengineers/SANKHYA: the source, the CSVs, the CI runs.",
        "docs/REPORT.pdf: this evidence as one document with a 25-question jury Q&A, every "
        "answer naming its file.",
        "docs/BENCHMARKS.md: every suite, every failure named, generated from the CSVs.",
        "docs/VIDEO.md: the four-minute walkthrough, command by command.",
        "sankhya-solver.vercel.app: the problem, the solver, the evidence, how to run it, the "
        "team.",
        "Clone it, build it, run scripts/reproduce.sh, and report any number that does not "
        "reproduce."], 16, NAVY)
    assert len(d.prs.slides) == SLIDES, len(d.prs.slides)
    d.save(target)


if __name__ == "__main__":
    import make_report
    make_report.gather()
    out = REPO_ROOT / "docs" / "SANKHYA-deck.pptx"
    write(make_report.F, out)
    print(f"wrote {out.relative_to(REPO_ROOT)}, {SLIDES} slides")
